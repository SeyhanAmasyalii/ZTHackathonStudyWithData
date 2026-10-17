# Mercek — hackathon sunumu, Ziraat Teknoloji sunum şablonuyla (.pptx, 16:9).
# Çalıştırma: python3 build_deck.py
# Şablon: sablon/ziraat-sunum-sablonu-2025.pptx (logo ve asıl slayt buradan gelir).
# Şablon kuralları: beyaz zemin, Calibri, başlık 32 punto kalın, metin 14-20 punto,
# kırmızı ve gri tonları, sayfada en fazla 5 renk, en fazla 5 madde.
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
TEMPLATE = HERE / "sablon" / "ziraat-sunum-sablonu-2025.pptx"
OUT = HERE / "Mercek_AI_Overflow_Sunum.pptx"
MOCK = HERE / "mockups"
ICON = HERE / "assets" / "icons"
LOGO = HERE / "assets" / "logo"

# Şablondaki renkler
RED, RED2, RED3 = RGBColor(0xE1, 0x05, 0x14), RGBColor(0xB9, 0x0F, 0x0A), RGBColor(0x87, 0x0F, 0x00)
G1, G2, G3 = RGBColor(0xD2, 0xDC, 0xE1), RGBColor(0xB4, 0xBE, 0xC3), RGBColor(0x91, 0x96, 0x9B)
INK, WHITE, SOFT = RGBColor(0x1E, 0x1E, 0x1E), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF2, 0xF5, 0xF6)
FONT = "Calibri"

prs = Presentation(TEMPLATE)
# Şablondaki örnek slaytları kaldır (asıl slayt, düzenler ve logo kalır)
sld_ids = prs.slides._sldIdLst
for sid in list(sld_ids):
    prs.part.drop_rel(sid.get(qn("r:id")))
    sld_ids.remove(sid)
BLANK = next(l for l in prs.slide_layouts if l.name == "Boş")
page = {"n": 0}


# ---------- yardımcılar ----------
def new_slide(title=None, sub=None):
    s = prs.slides.add_slide(BLANK)
    page["n"] += 1
    if title:
        text(s, 0.15, 0.1, 11.0, 0.62, [(title, {"size": 32, "bold": True})])
        wordmark(s, 11.55, 0.17, 0.6)
        if sub:
            text(s, 0.17, 0.72, 12.5, 0.4, [(sub, {"size": 16, "color": G3})])
    if page["n"] > 1:
        text(s, 12.35, 6.98, 0.75, 0.3, [(str(page["n"]), {"size": 11, "color": G3})], align=PP_ALIGN.RIGHT)
    return s


def _run(p, t, o):
    r = p.add_run()
    r.text = t
    f = r.font
    f.name = FONT
    f.size = Pt(o.get("size", 16))
    f.bold = o.get("bold", False)
    f.italic = o.get("italic", False)
    f.color.rgb = o.get("color", INK)


def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, space=4, name=None):
    """paras: list of (text, opts) — her biri bir paragraf; ya da [(t,o),(t,o)] listesi = tek paragrafta birden çok run."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = para if isinstance(para, list) else [para]
        for t, o in runs:
            _run(p, t, o)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "228600")
            pPr.set("indent", "-228600")
            bu_clr = pPr.makeelement(qn("a:buClr"), {})
            srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": "E10514"})
            bu_clr.append(srgb)
            pPr.append(bu_clr)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
    return tb


def box(s, x, y, w, h, fill=SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.text = ""
    return sh


def circle(s, x, y, d, fill, label=None, size=16, icon=None):
    c = box(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    if icon:
        p = d * 0.25
        s.shapes.add_picture(str(ICON / icon), Inches(x + p), Inches(y + p), Inches(d - 2 * p), Inches(d - 2 * p))
    if label:
        text(s, x, y, d, d, [(label, {"size": size, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return c


def pill(s, x, y, w, h, label, fill, color=WHITE, size=11):
    box(s, x, y, w, h, fill=fill, radius=0.5)
    text(s, x, y, w, h, [(label, {"size": size, "bold": True, "color": color})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def line(s, x1, y1, x2, y2, color=G3):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))



def wordmark(s, x, y, size=1.0):
    """Köşe logosu: sade mercek + MERCEK yazısı."""
    d = 0.64 * size
    s.shapes.add_picture(str(LOGO / "mercek-logo-kucuk.png"), Inches(x), Inches(y), Inches(d), Inches(d))
    text(s, x + d + 0.12 * size, y, 3 * size, d, [("MERCEK", {"size": int(26 * size), "bold": True})], anchor=MSO_ANCHOR.MIDDLE)


def logo(s, x, y, w):
    """Büyük logo: saplı mercek, log satırı odakta."""
    s.shapes.add_picture(str(LOGO / "mercek-logo.png"), Inches(x), Inches(y), Inches(w), Inches(w))

def draft_note(s, x, y):
    text(s, x, y, 7, 0.28, [("Taslak görsel — gerçek ekran görüntüsüyle değiştirilecek", {"size": 11, "italic": True, "color": G3})])


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ---------- 1. Kapak ----------
s = new_slide()
text(s, 1.0, 1.55, 6.8, 0.4, [("ZİRAAT TEKNOLOJİ YAPAY ZEKÂ HACKATHONU · TEMA 2", {"size": 14, "bold": True, "color": RED})])
text(s, 1.0, 2.0, 6.8, 1.3, [("MERCEK", {"size": 72, "bold": True})])
box(s, 1.05, 3.35, 1.2, 0.07, fill=RED, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.0, 3.6, 6.8, 0.6, [("Soruna odaklanır, çözümü netleştirir.", {"size": 26, "italic": True, "color": RED2})])
text(s, 1.0, 4.75, 6.8, 0.9, [("AI Overflow Takımı", {"size": 20, "bold": True}), ("17 Ekim 2026", {"size": 16, "color": G3})], space=4)
logo(s, 7.75, 1.0, 4.6)
notes(s, "0:00–0:10 · Merhaba, biz AI Overflow. Projemizin adı Mercek: soruna odaklanır, çözümü netleştirir.")

# ---------- 2. Takım ----------
s = new_slide("AI Overflow Takımı")
for i in range(5):
    x, y, w, h = 0.55 + i * 2.5, 1.45, 2.3, 4.6
    box(s, x, y, w, h)
    c = box(s, x + w / 2 - 0.65, y + 0.4, 1.3, 1.3, fill=WHITE, line=RED, shape=MSO_SHAPE.OVAL)
    s.shapes.add_picture(str(ICON / "user-g.png"), Inches(x + w / 2 - 0.3), Inches(y + 0.75), Inches(0.6), Inches(0.6))
    text(s, x + 0.15, y + 1.95, w - 0.3, 2.2, [("Ad Soyad", {"size": 18, "bold": True}), ("Birim", {"size": 14, "color": G3}),
                                               ("Projedeki rolü", {"size": 14})], align=PP_ALIGN.CENTER, space=6)
text(s, 0.55, 6.25, 12, 0.35, [("Sorumluluklar 19 maddelik listeden herkesin kendi seçimiyle dağıtıldı.", {"size": 14, "color": G3})])
notes(s, "0:10–0:20 · Takımı tek cümleyle tanıtın. (Ad, birim ve rolleri doldurun; fotoğraf eklenecekse daire içine yerleştirin.)")

# ---------- 3. Problem ----------
s = new_slide("Hata mı, Kural mı? Bilinmeyince Çağrı Açılıyor")
box(s, 0.45, 1.3, 3.55, 5.25, fill=RED2)
s.shapes.add_picture(str(ICON / "quote-w.png"), Inches(0.8), Inches(1.65), Inches(0.5), Inches(0.5))
text(s, 0.8, 2.35, 2.9, 2.4, [("“Ekran ‘kullandırılamaz’ diyor. Sistem mi bozuk?”", {"size": 24, "bold": True, "italic": True, "color": WHITE})])
text(s, 0.8, 5.6, 2.9, 0.7, [("Elif, şube müşteri temsilcisi · KRDTKULL", {"size": 14, "color": WHITE})], anchor=MSO_ANCHOR.BOTTOM)
steps = ["Temsilci uyarıyı görür", "Çağrı açar", "L1 Kibana'da log arar", "L2 kurumu ve kodu inceler", "Yeni talep bu sorunlardan habersiz planlanır"]
for i, t in enumerate(steps):
    x = 4.3 + i * 1.72
    box(s, x, 1.3, 1.58, 1.95)
    circle(s, x + 0.15, 1.45, 0.45, RED, str(i + 1), 14)
    text(s, x + 0.15, 2.0, 1.3, 1.2, [(t, {"size": 14})])
text(s, 4.3, 3.4, 8.5, 0.6, [[("Bedeli: ", {"size": 14, "bold": True, "color": RED2}),
                             ("işlem yapamayan müşteri · gereksiz çağrılar · aynı kurum sorunu için onlarca kayıt · veriden kopuk planlanan talepler", {"size": 14, "color": RED2})]])
pains = [("eye-w.png", "Hata mı kural mı belli değil", "İş kuralı reddi ekranda teknik hatadan ayrılamıyor."),
         ("users-w.png", "Aynı sorun, onlarca çağrı", "Kurum kaynaklı sorunda her arama ayrı kayıt; çözüm kişilerde."),
         ("unlink-w.png", "Talepler veriden kopuk", "Efor toplantıda; erişim ve kapasite sorunları geç fark ediliyor.")]
for i, (ic, h1, d) in enumerate(pains):
    x = 4.3 + i * 2.88
    circle(s, x, 4.25, 0.6, G3, icon=ic)
    text(s, x, 5.0, 2.7, 1.5, [(h1, {"size": 16, "bold": True}), (d, {"size": 14, "color": G3})])
notes(s, "0:20–0:40 · Şube personeli işlem yapamadığında hata mı kural mı bilemiyor, çağrı açıyor. L1 ve L2 Kibana'da log arıyor, aynı kurum sorunu için onlarca çağrı açılıyor, bilinen çözümler kişilerde kalıyor. Yeni talepler de bu sorunlardan habersiz planlanıyor.")

# ---------- 4. Çözüm ----------
s = new_slide("Mercek: Sorundan Plana Tek Akış")
text(s, 0.45, 1.2, 12.4, 0.8, [("Banka çalışanlarının ekranlarda yaşadığı sorunu loglardan, çağrı geçmişinden ve koddan açıklar; çağrıları doğru servise yönlendirip birleştirir ve iş birimi taleplerini servis bazlı efor ve takvimle plana dönüştürür.", {"size": 18})])
mods = [("headset-w.png", "Akıllı Problem Çözüm ve Çağrı Yönetimi",
         ["Ekran bağlamıyla logu bulur, iş kuralını hatadan ayırır", "Çağrıyı servise atar, mükerrerleri ana kayda bağlar", "Geçmiş çözümü bulur; yoksa duyuruya ve koda bakar"]),
        ("clip-w.png", "Akıllı Talep Analizi, Efor ve Takvim",
         ["Talebi bileşenlere ayırır, ekranları ve servisleri bulur", "Benzer taleplerden adam-gün, servis takviminden plan", "Çağrı ve müşteri verisiyle kapsamı düzeltir"])]
for i, (ic, head, items) in enumerate(mods):
    x, y, w = 0.45 + i * 6.65, 2.2, 5.75
    box(s, x, y, w, 3.25)
    box(s, x, y, w, 0.85, fill=RED if i == 0 else RED2, radius=0.12)
    circle(s, x + 0.2, y + 0.13, 0.6, RED3, icon=ic)
    text(s, x + 0.95, y, w - 1.1, 0.85, [(head, {"size": 18, "bold": True, "color": WHITE})], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, y + 1.1, w - 0.6, 2.0, [(t, {"size": 16}) for t in items], bullets=True, space=8)
box(s, 6.27, 3.3, 0.8, 1.0, fill=RED3, radius=0.15)
text(s, 6.27, 3.3, 0.8, 1.0, [("veri", {"size": 11, "bold": True, "color": WHITE}), ("→", {"size": 11, "bold": True, "color": WHITE}), ("kapsam", {"size": 11, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=0)
text(s, 0.45, 5.75, 1.8, 0.55, [("Kullanıcılar", {"size": 16, "bold": True})], anchor=MSO_ANCHOR.MIDDLE)
for i, t in enumerate(["Şube, çağrı merkezi", "L2, platform", "İş birimleri", "İş analisti"]):
    pill(s, 2.3 + i * 2.6, 5.75, 2.4, 0.55, t, G1, INK, 14)
notes(s, "0:40–0:50 · Tek cümle çözüm. İki modül ve köprü: loglar, çağrı ve müşteri verisi talep kapsamına giriyor. Mercek'i müşteri değil banka çalışanları kullanıyor.")

# ---------- 5. Farklar ----------
s = new_slide("Dört Farkımız", "Bugünkü araçlar sorunu kaydeder; Mercek açıklar, kaynağına bağlar ve eyleme dönüştürür.")
diffs = [("bolt-w.png", RED, "Bağlamı bilir", "Kullanıcının görevini ve kendi logunu bilir; ekran adı söylenmese de belirtiden hatayı anlamca bulur."),
         ("code-w.png", RED2, "Dört kaynağı birleştirir", "Log (ne oluyor), geçmiş çağrı (nasıl çözüldü), hata türü ve hazır çözüm araması tek cevapta."),
         ("sync-w.png", RED3, "Sorundan eyleme", "Çözümün ekran akışını açar; L2 ekibine etkilenen müşteri listesini KVKK uyumlu verir."),
         ("shield-w.png", G3, "Güvenilir AI", "Hesabı kod yapar, AI açıklar. %75 eşiğin altında öneri yok; her rakam araçtan gelir.")]
for i, (ic, col, head, d) in enumerate(diffs):
    x, y, w, h = 0.45 + (i % 2) * 6.3, 1.4 + (i // 2) * 2.6, 6.05, 2.35
    box(s, x, y, w, h)
    circle(s, x + 0.35, y + 0.4, 0.85, col, icon=ic)
    text(s, x + 1.5, y + 0.35, w - 1.8, h - 0.5, [(head, {"size": 20, "bold": True}), (d, {"size": 16, "color": INK})], space=6)
notes(s, "0:50–0:58 · Mercek çağrıyı kaydetmekle kalmıyor: açıklıyor, birleştiriyor, koda bağlıyor ve talebi gerekçelendiriyor. AI'ın rolü sınırlı ve denetlenebilir.")

# ---------- 6. Mimari ----------
s = new_slide("Mimari: Foundry Agent + Hibrit RAG")


def arch(x, y, w, h, head, sub, fill, color=INK, line_c=None, hs=15, ss=12):
    box(s, x, y, w, h, fill=fill, line=line_c)
    text(s, x + 0.08, y + 0.05, w - 0.16, h - 0.1, [(head, {"size": hs, "bold": True, "color": color}), (sub, {"size": ss, "color": color})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space=2)


arch(0.45, 2.35, 1.6, 2.6, "Kullanıcılar", "Şube\nÇağrı merkezi\nL2 destek", SOFT)
arch(2.3, 1.35, 2.3, 1.75, "App Service · web", "Next.js 16\nsohbet · Akışı izle\nL2 müşteri tablosu", RED, WHITE)
arch(2.3, 4.15, 2.3, 1.75, "App Service · api", "FastAPI · araçları çalıştırır\nhesap, eşik, yetki\nKVKK, erişim kaydı", RED2, WHITE)
# Foundry agent + veri çerçevesi
box(s, 4.85, 1.2, 4.1, 4.75, fill=WHITE, line=RED)
text(s, 5.0, 1.24, 3.8, 0.4, [[("AGENT + HİBRİT RAG  ", {"size": 13, "bold": True, "color": RED}),
                              ("araç seçer · kod çalıştırır", {"size": 11, "color": G3})]], anchor=MSO_ANCHOR.MIDDLE)
arch(7.55, 1.7, 1.25, 3.2, "Foundry Agent", "mercek-cagri\ngpt-5.5\n\n4 araç\n\nizleme:\nApp Insights", RED, WHITE, hs=13, ss=11)
arch(5.0, 1.7, 2.3, 1.6, "Azure AI Search", "hata türleri (107)\nhazır çözüm + çağrı\nnotu (37) · %75 eşik", SOFT, line_c=RED, hs=14, ss=11)
arch(5.0, 3.45, 2.3, 1.45, "Azure SQL", "log · çağrı · müşteri\nkonuşma geçmişi\nerişim kaydı", SOFT, line_c=G2, hs=14, ss=11)
text(s, 5.0, 5.0, 3.8, 0.9, [("Agent hangi aracı çağıracağına karar verir; sorguyu backend çalıştırır. Rakam ve eşik koddan gelir, model yalnızca açıklar.",
                              {"size": 11, "color": INK})], anchor=MSO_ANCHOR.MIDDLE)
line(s, 2.05, 3.0, 2.3, 2.25)
line(s, 3.45, 3.1, 3.45, 4.15)
line(s, 4.6, 4.7, 5.0, 2.5)
line(s, 4.6, 5.0, 5.0, 4.17)
pill(s, 5.35, 3.22, 1.6, 0.3, "▲ indexer · 5 dk", RED3, WHITE, 10)
box(s, 0.45, 6.1, 8.35, 0.65)
s.shapes.add_picture(str(ICON / "shield-g.png"), Inches(0.6), Inches(6.22), Inches(0.42), Inches(0.42))
text(s, 1.15, 6.1, 7.6, 0.65, [[("Parolasız: ", {"size": 12, "bold": True}), ("Managed Identity ile SQL, AI Search ve Foundry    ", {"size": 12}),
                               ("GHAS: ", {"size": 12, "bold": True}), ("CodeQL · Dependabot · secret scanning", {"size": 12})]], anchor=MSO_ANCHOR.MIDDLE)
box(s, 9.15, 1.35, 3.75, 2.55, fill=RED)
text(s, 9.35, 1.45, 3.4, 2.4, [("AI ne yapar (Foundry)", {"size": 16, "bold": True, "color": WHITE})] +
     [(t, {"size": 14, "color": WHITE}) for t in ["Belirtiyi anlar, ekranı bulur", "Hangi aracı çağıracağına karar verir", "Hazır çözümü kullanıcıya uyarlar", "Sistemsel riski açıklar"]], space=3)
box(s, 9.15, 4.1, 3.75, 2.65, fill=G3)
text(s, 9.35, 4.2, 3.4, 2.5, [("Kod ne yapar (kurallar)", {"size": 16, "bold": True, "color": WHITE})] +
     [(t, {"size": 14, "color": WHITE}) for t in ["Log sayımı, zaman aralığı", "Eşleşme yüzdesi, %75 eşik", "L2 yetkisi, telefon maskesi", "Akış bağlantısı seçimi", "Erişim kaydı (audit)"]], space=3)
notes(s, "0:58–1:05 · Mercek, Foundry Agent Service'te tanımlı mercek-cagri agent'ını kullanır. Agent dört araçtan hangisini çağıracağına karar verir: kibana_sorgula ve cagri_ara Azure SQL'deki log ve çağrı verisini, hata_ara ve benzer_cozum_ara AI Search'teki hata türü ve hazır çözüm indekslerini sorgular. Araçları backend çalıştırır; rakam, eşleşme yüzdesi (%75 eşik) ve yetki kodda. AI Search indeksleri SQL'den indexer ile 5 dakikada bir kendiliğinden güncellenir. Konuşma geçmişi ve KVKK erişim kaydı SQL'de; agent'ın her adımı Foundry'de App Insights ile izlenir. Tüm servis bağlantıları Managed Identity ile, parolasız.")

# ---------- 7. Kullanım senaryoları ----------
s = new_slide("Altı Kullanım Senaryosu")
# live: True = Mercek'te çalışıyor ve canlı gösterilir; False = planlanan
scen = [("A", "Hata mı, kural mı?", "KRDTKULL · Elif (şube)", "Taksit gelirin %50'sini aşıyor: iş kuralı açıklanır", "Çağrı açılmadan çözülür", True),
        ("B", "Yaygın sorun, bilinen çözüm", "ODMTTLMT · GAZ02 · Elif, Zeynep, Kerem", "Kurum zaman aşımı, SOL-031 ve adım adım akış", "6 benzer çağrı · 23 kayıt", True),
        ("C", "Kural yanlış kapatılmış", "ODMTTLMT · İSKİ 13 hane · Elif → Kerem", "158 kayıt, tek kurumda; 'kullanıcı hatası' temkinli", "Sistemsel olabilir uyarısı", True),
        ("D", "Etkilenen müşteriler", "Fatura talimatı · Kerem (L2)", "Bilgilendirme listesi; izinsiz telefon maskeli", "31 müşteri · 11 aranabilir", True),
        ("E", "Talepten plana", "Talep analizi · Fatma → Ayşe", "Servis, adam-gün, takvim; veriyle kapsam", "Modül 2 · sırada", False),
        ("F", "Günlük durum özeti", "Mercek sohbet · Kerem", "Kibana bilgisi gerektirmeden sohbetle özet", "Son 30 dk: 23 hata", True)]
for i, (L, head, who, what, res, live) in enumerate(scen):
    x, y, w, h = 0.45 + (i % 3) * 4.2, 1.2 + (i // 3) * 2.8, 3.95, 2.6
    col = RED if live else G3
    box(s, x, y, w, h)
    circle(s, x + 0.25, y + 0.25, 0.62, col, L, 20)
    text(s, x + 1.02, y + 0.2, w - 1.15, 0.75, [(head, {"size": 17, "bold": True}), (who, {"size": 12, "color": G3})], anchor=MSO_ANCHOR.MIDDLE, space=1)
    text(s, x + 0.25, y + 1.07, w - 0.5, 0.7, [(what, {"size": 14})])
    text(s, x + 0.25, y + 1.85, w - 1.55, 0.55, [(res, {"size": 14, "bold": True, "color": RED2 if live else INK})], anchor=MSO_ANCHOR.MIDDLE)
    pill(s, x + w - 1.3, y + 1.93, 1.1, 0.38, "CANLI" if live else "PLAN", col)
notes(s, "1:05–1:20 · 'CANLI' etiketli beşi Mercek'te çalışıyor ve birazdan gösterilecek; 'PLAN' Modül 2 (talep analizi), sıradaki iş. A: KRDTKULL 120218, taksit gelirin %50'sini aşıyor; iş kuralı. B: Başkentgaz kurum zaman aşımı, SOL-031 ile kapanmış 6 çağrı, cevabın altında 'Akışı izle'. C: İSKİ 13 hane, 158 kayıt; geçmişte 'kullanıcı hatası' diye kapatılmış ama tek kurumda yoğunlaşıyor, Mercek temkinli. D: L2 için etkilenen müşteri listesi, KVKK'ya uygun. F: Kerem 'Son durum ne?' diye sorar.")

# ---------- 8. Canlı demo ----------
s = new_slide("Canlı Demo: Dört Sahne", "Her cevapta araç adımları Foundry izlemesinde görünür; rakamlar koddan gelir.")
scenes = [("~45 sn", "headset-w.png", RED, "B · Bilinen çözüm", "Elif",
           ["“ODMTTLMT ekranında hata alıyorum”", "Kendi logunda 7 hata · GAZ02 110731", "SOL-031 ile kapanmış 6 çağrı", "▶ Akışı izle: adım adım ekran turu"]),
          ("~40 sn", "eye-w.png", RED2, "C · Belirtiyle arama", "Elif → Kerem",
           ["“abone numarasını sistem kabul etmiyor”", "AI Search: ODMTTLMT 110412 · İSKİ 13 hane", "7 günde 158 kayıt, tek kurumda", "‘Kullanıcı hatası’ çözümüne temkin"]),
          ("~40 sn", "users-w.png", RED3, "D · Etkilenen müşteriler", "Kerem (L2)",
           ["“Aktif talimat zaten var… kaç müşteri?”", "Etkilenen müşteriler (31) düğmesi", "11 aranabilir; izinsiz telefon maskeli", "Görüntüleme erişim kaydına yazılır"]),
          ("~35 sn", "shield-w.png", G3, "Güvenilir AI", "Kerem (L2)",
           ["“Input string was not in a correct format”", "Hazır çözüm araması: %75 eşiğin altında", "Uydurmaz: ‘hazır çözüm yok, ekibe ilet’", "Foundry izlemesinde araç adımları"])]
for i, (t, ic, col, head, who, steps) in enumerate(scenes):
    x, y, w, h = 0.45 + (i % 2) * 6.3, 1.3 + (i // 2) * 2.75, 6.1, 2.55
    box(s, x, y, w, h)
    circle(s, x + 0.25, y + 0.25, 0.75, col, icon=ic)
    text(s, x + 1.2, y + 0.22, w - 2.5, 0.8, [(head, {"size": 18, "bold": True}), (who, {"size": 12, "color": G3})], anchor=MSO_ANCHOR.MIDDLE, space=1)
    pill(s, x + w - 1.2, y + 0.32, 0.95, 0.34, t, col)
    text(s, x + 0.3, y + 1.15, w - 0.6, 1.35, [(st, {"size": 14, "bold": j == 0, "italic": j == 0}) for j, st in enumerate(steps)], space=2)
notes(s, "1:20–4:15 · CANLI DEMO (~3 dk). B: Elif → 'ODMTTLMT ekranında hata alıyorum' → 7 hata, GAZ02 110731, SOL-031 (6 çağrı) → 'Akışı izle'ye tıklayın; aynı sohbette 'Peki nasıl çözerim?' (önceki konuyu hatırlar). C: Yeni sohbet → 'abone numarasını sistem kabul etmiyor' → AI Search hata türünü bulur, İSKİ 13 hane, SOL-005 temkinli. D: Kerem → K8 sorusu → 'Etkilenen müşteriler (31)' → tablo, 11 aranabilir, maskeli telefonlar. Güvenilir AI: Kerem → K5 → hazır çözüm yok. Sonra Foundry → Aracılar → mercek-cagri → İzlemeler: araç adımları. Yedek: sorun çıkarsa LLM_MODE=chat; ekran kaydına geçin.")

# ---------- 9. İş değeri ----------
s = new_slide("İş Değeri: Önce ve Sonra")
for i, (big, lbl, col) in enumerate([("Çağrı hiç açılmıyor", "İş kuralı reddinde (A)", RED), ("Saniyeler içinde çözüm", "Bilinen sorunda, kaynağıyla (B)", RED2), ("Tek tıkla liste", "Etkilenen müşteri, KVKK uyumlu (D)", RED3)]):
    x = 0.45 + i * 4.2
    box(s, x, 1.2, 3.95, 1.5)
    text(s, x + 0.3, 1.2, 3.4, 1.5, [(big, {"size": 22, "bold": True, "color": col}), (lbl, {"size": 14, "color": G3})], anchor=MSO_ANCHOR.MIDDLE, space=2)
rows = [("Ölçü", "Bugün", "Mercek ile"),
        ("İş kuralı reddi", "Çağrı açılır, L1 kuralı açıklar", "Mercek kuralı ve nedenini açıklar; çağrı açılmaz"),
        ("Yaygın kurum sorunu", "Her arama ayrı çağrı", "Kurum kaynaklı olduğu, kaç kayıtta görüldüğüyle söylenir"),
        ("Bilinen sorun", "L2 log arar, çözüm kişide", "Hazır çözüm ve ekran akışı; %75 eşiğin altında öneri yok"),
        ("Yanlış kapatılan kural", "Kullanıcı hatası diye kapanır, tekrarlar", "Tek kurumda yoğunlaşma işaretlenir, temkinli sunulur"),
        ("Etkilenen müşteri", "Liste elle çıkarılır, izin bakılmaz", "L2 için tek tık; izinsiz numara maskeli, erişim kayıtlı")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.45), Inches(2.95), Inches(12.4), Inches(3.3)).table
for j, wdt in enumerate([2.9, 4.3, 5.2]):
    tbl.columns[j].width = Inches(wdt)
tblPr = tbl._tbl.tblPr
for k in list(tblPr):
    if k.tag == qn("a:tableStyleId"):
        tblPr.remove(k)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED2 if r == 0 else (SOFT if r % 2 == 0 else WHITE)
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.text = ""
        _run(tf.paragraphs[0], val, {"size": 14, "bold": r == 0 or c == 0, "color": WHITE if r == 0 else INK})
text(s, 0.45, 6.4, 12.4, 0.3, [("Rakamlar sentetik veriye dayanır; her rakam Mercek'in araçlarından gelir ve Foundry izlemesinde doğrulanabilir.", {"size": 12, "italic": True, "color": G3})])
notes(s, "4:15–4:45 · Önce/sonra. Demodaki sayılarla bağlayın: 7 hata ve SOL-031 saniyeler içinde; İSKİ tarafında 158 kayıt tek kurumda; 31 etkilenen müşteri, 11 aranabilir. Veri sentetik.")

# ---------- 10. Neredeyiz ----------
s = new_slide("Neredeyiz, Sırada Ne Var")
cols = [("check-w.png", RED, "Tamamlanan", ["M1 sohbet: log, çağrı, hata türü ve hazır çözüm araçları", "Foundry agent (mercek-cagri) ve izleme", "L2 etkilenen müşteri listesi, erişim kaydı", "Azure: 2 App Service, SQL, AI Search, Managed Identity"]),
        ("warn-w.png", G3, "Sınırlamalar", ["Veri tamamen sentetik", "Gerçek log, çağrı ve kod entegrasyonu yok", "Çözüm önerisi; karar insanda", "Arama ve SMS yok, liste bilgilendirme için"]),
        ("route-w.png", RED3, "Sırada", ["KVKK: modele gitmeden metin taraması", "Duyuru ve işlenmemiş e-posta araması (İSKİ)", "Modül 2: talep analizi, efor, takvim", "Pilot: şube kanalı, ODMTTLMT ve KRDTKULL"])]
for i, (ic, col, head, items) in enumerate(cols):
    x, w = 0.45 + i * 4.2, 3.95
    box(s, x, 1.2, w, 5.4)
    circle(s, x + 0.3, 1.45, 0.75, col, icon=ic)
    text(s, x + 1.2, 1.45, w - 1.4, 0.75, [(head, {"size": 20, "bold": True})], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.3, 2.5, w - 0.6, 3.9, [(t, {"size": 15}) for t in items], bullets=True, space=10)
notes(s, "4:45–4:55 · Tamamlanan: Modül 1 sohbeti dört araçla çalışıyor, Foundry agent ve izleme açık, L2 için KVKK'ya uygun müşteri listesi var, tüm bağlantılar parolasız. Sırada: modele gitmeden kişisel veri taraması, kurum e-postalarında kök neden araması ve Modül 2. Soru gelirse: AI yanlışsa → %75 eşik, kaynak, koddan rakam, insan onayı. Maliyet → tek agent, kısa bağlam, araç sonuçları özet.")

# ---------- 11. Ekler ve kapanış ----------
s = new_slide("Ekler")
for i, (ic, h1, v) in enumerate([("github-w.png", "GitHub repository", "[repo bağlantısı]"), ("globe-w.png", "Web uygulaması", "[mercek-web adresi]"), ("video-w.png", "Senaryo videoları", "[video bağlantısı]")]):
    y = 1.5 + i * 1.2
    circle(s, 0.6, y, 0.8, RED2, icon=ic)
    text(s, 1.65, y - 0.02, 5.5, 0.85, [(h1, {"size": 18, "bold": True}), (v, {"size": 16, "color": RED})], anchor=MSO_ANCHOR.MIDDLE, space=2)
logo(s, 8.45, 1.25, 3.3)
text(s, 7.6, 4.85, 5.0, 1.0, [("MERCEK", {"size": 32, "bold": True}), ("Soruna odaklanır, çözümü netleştirir.", {"size": 18, "italic": True, "color": RED2})], align=PP_ALIGN.CENTER, space=4)
text(s, 0.6, 5.4, 12, 0.5, [("Teşekkürler · AI Overflow", {"size": 20, "bold": True, "color": RED2})])
notes(s, "4:55–5:00 · Sloganla kapatın. Bağlantıları sunumdan önce doldurun; şifre, anahtar veya gerçek müşteri verisi görünmediğini kontrol edin.")

prs.save(OUT)
print("OK", OUT)
